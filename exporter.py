# -*- coding: utf-8 -*-
"""
Excel 导出：把分类 JSON 导出为多 sheet Excel 工作簿。
sheet：主题汇总 / 问题清单 / 全部帖子 / 每日趋势 / 周环比 / 热点叙事线
"""
import argparse
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="993C1D")
THIN = Side(style="thin", color="E5E3DA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_sheet(ws, widths, freeze="A2", num_cols=()):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = Font(color="FFFFFF", bold=True, size=11)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = freeze
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.border = BORDER
            if c.column in num_cols:
                c.alignment = Alignment(horizontal="right")
                c.number_format = "#,##0"


def export_xlsx(data_path, out_path):
    """导出多 sheet Excel 工作簿"""
    d = json.load(open(data_path, encoding="utf-8"))
    cats, meta = d["categories"], d["meta"]
    wb = Workbook()

    # ---- 主题汇总 ----
    ws = wb.active
    ws.title = "主题汇总"
    ws.append(["主题", "类别", "帖子数", "问题主题数", "总回复", "平均回复", "总浏览"])
    for c in cats:
        ws.append([c["name"], "补充" if c.get("supplementary") else "限定",
                   c["count"], len(c.get("themes", [])), c["total_replies"],
                   c["avg_replies"], c["total_views"]])
    ws.append([])
    ws.append(["统计区间", meta["date_range"]])
    ws.append(["周期内活跃帖", meta["total_raw_posts"]])
    ws.append(["起点相关帖(闸门内)", meta["gated_total"]])
    ws.append(["生成时间", meta["generated_at"]])
    style_sheet(ws, [30, 8, 10, 12, 12, 10, 14], num_cols=(3, 4, 5, 6, 7))

    # ---- 问题清单 ----
    ws = wb.create_sheet("问题清单")
    ws.append(["主题", "类别", "问题主题", "涉及篇数", "合计回复", "代表帖TOP5"])
    for c in cats:
        tag = "补充" if c.get("supplementary") else "限定"
        for t in c.get("themes", []):
            tops = " | ".join(f"{p['title']}（{p['replies']}回复）" for p in t["top_posts"])
            ws.append([c["name"], tag, t["theme"], t["count"], t["total_replies"], tops])
    style_sheet(ws, [28, 8, 30, 10, 12, 95], num_cols=(4, 5))
    for r in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        r[0].alignment = Alignment(wrap_text=True, vertical="top")

    # ---- 全部帖子 ----
    ws = wb.create_sheet("全部帖子")
    ws.append(["主题", "类别", "问题主题", "帖子标题", "回复", "浏览", "发帖日", "作者", "链接"])
    for c in cats:
        tag = "补充" if c.get("supplementary") else "限定"
        for t in c.get("themes", []):
            for p in t.get("all_posts", t["top_posts"]):
                ws.append([c["name"], tag, t["theme"], p["title"], p["replies"],
                           p.get("views", 0), p["date"], p.get("author") or "", p["url"]])
    style_sheet(ws, [26, 8, 27, 52, 8, 10, 12, 14, 42], num_cols=(5, 6))
    for r in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        if r[0].value:
            r[0].font = Font(color="378ADD", underline="single")

    # ---- 每日趋势 ----
    ws = wb.create_sheet("每日趋势")
    ws.append(["日期", "起点相关帖", "AI治理判定"])
    for k, v in meta["daily"].items():
        ws.append([k, v, meta["ai_daily"].get(k, 0)])
    style_sheet(ws, [14, 38, 45], num_cols=(2, 3))

    # ---- 周环比 ----
    wc = meta.get("week_compare")
    if wc:
        ws = wb.create_sheet("周环比")
        ws.append(["主题", "类别", "上周帖数", "本周帖数", "帖数变化率",
                   "上周回复", "本周回复", "回复变化率"])
        for r in wc["categories"]:
            rc = (r["cur_count"] - r["prev_count"]) / r["prev_count"] if r["prev_count"] else None
            rr = (r["cur_replies"] - r["prev_replies"]) / r["prev_replies"] if r["prev_replies"] else None
            ws.append([r["name"], "补充" if r["supplementary"] else "限定",
                       r["prev_count"], r["cur_count"],
                       f"{rc:+.0%}" if rc is not None else "新增",
                       r["prev_replies"], r["cur_replies"],
                       f"{rr:+.0%}" if rr is not None else "新增"])
        t = wc["total"]
        ws.append([])
        ws.append(["合计", "", t["prev_gated"], t["cur_gated"],
                   f"{(t['cur_gated']-t['prev_gated'])/t['prev_gated']:+.0%}" if t["prev_gated"] else "",
                   t["prev_replies"], t["cur_replies"],
                   f"{(t['cur_replies']-t['prev_replies'])/t['prev_replies']:+.0%}" if t["prev_replies"] else ""])
        ws.append([])
        ws.append([f"本期 {wc['cur_range']} vs 上期 {wc['prev_range']}（同口径分别分类）"])
        style_sheet(ws, [31, 29, 37, 41, 43, 47, 49, 53], num_cols=(3, 4, 6, 7))

    # ---- 热点叙事线 ----
    st = meta.get("storylines")
    if st and st.get("lines"):
        ws = wb.create_sheet("热点叙事线")
        ws.append(["叙事线", "引爆点", "跨线交点", "落点", "序", "节点日", "帖子标题", "回复", "链接"])
        for l in st["lines"]:
            for i, n in enumerate(l["nodes"]):
                ws.append([l["name"], l["trigger"], l["link"], l["close"],
                           i + 1, n["day"], n["title"], n.get("replies", 0), n["url"]])
        style_sheet(ws, [55, 57, 58, 56, 51, 63, 66, 68, 65], num_cols=(8,))

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    print(f"已导出: {out_path} | sheet: {wb.sheetnames}")
    return out_path


def main():
    ap = argparse.ArgumentParser(description="导出 Excel 工作簿")
    ap.add_argument("--data", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    export_xlsx(args.data, args.out)


if __name__ == "__main__":
    main()
